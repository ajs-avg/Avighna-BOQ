"""Fit-out Bill of Quantities workbook.

Reproduces the KR-style BOQ used on the Avighna office / sales-centre job
('BOQ - Avigna group ... R00.xlsx'): a COVER, one tab per division
(DIV-01-PRELIMINARIES, DIV-09-FINISHES with joinery sections inside it,
DIV-08-DOORS, DIV-13-LIGHTING, FFE, ...), a SUMMARY with VAT, plus a
MEASUREMENT BACKUP tab that is the audit trail for every quantity.

Each division tab:
    ITEM | WORK DESCRIPTION | QTY | UNIT | RATE (ccy) | AMOUNT (ccy)
    section heading rows, one row per BOQ line (A, B, C ... per section),
    AMOUNT = QTY * RATE, "CARRIED TO COLLECTION" per section,
    a COLLECTION block and "TOTAL CARRIED TO SUMMARY".
"""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.export._style import (
    BORDER_BOTTOM_MEDIUM, BORDER_MEDIUM, BORDER_THIN, BRAND_RED, CENTER, FMT_MONEY,
    FMT_PCT, FMT_QTY, FMT_QTY_INT, GREY_FILL, LEFT, LEFT_BOTTOM, LEFT_TOP, PALE_FILL, RIGHT, WHITE,
    col, est_row_height, fill, font, header, merge, print_setup, put, qs, safe_sheet_name,
    set_heights, set_widths, style_cells, tab_color,
)
from app.models.schemas import (
    DIVISION_LABELS, BOQLine, BOQResult, BOQSection, FitoutItem, JobRecord, RoomRecord,
)

DIVISION_SEQ = ["preliminaries", "finishes", "joinery", "doors", "lighting", "ffe", "mep", "other"]

# division -> tab (joinery lives inside the finishes tab, like the reference)
TAB_FOR_DIVISION: Dict[str, str] = {
    "preliminaries": "DIV-01-PRELIMINARIES",
    "finishes": "DIV-09-FINISHES",
    "joinery": "DIV-09-FINISHES",
    "doors": "DIV-08-DOORS",
    "lighting": "DIV-13-LIGHTING",
    "ffe": "FFE",
    "mep": "DIV-15-MEP",
    "other": "OTHER",
}
TAB_SEQ = ["DIV-01-PRELIMINARIES", "DIV-09-FINISHES", "DIV-08-DOORS", "DIV-13-LIGHTING", "FFE",
           "DIV-15-MEP", "OTHER"]
# tab -> (page code, heading line 1, heading line 2, summary label)
TAB_META: Dict[str, Tuple[str, str, str, str]] = {
    "DIV-01-PRELIMINARIES": ("DIV01", "DIVISION 01", "PRELIMINARIES & GENERAL REQUIREMENTS", "DIV-01 - PRELIMINARIES"),
    "DIV-09-FINISHES": ("DIV09", "DIVISION 09", "FINISHES", "DIV-09 - FINISHES"),
    "DIV-08-DOORS": ("DIV08", "DIVISION 08", "DOORS", "DIV-08 - DOORS"),
    "DIV-13-LIGHTING": ("DIV13", "DIVISION 13", "LIGHTING", "DIV-13 - LIGHTING"),
    "FFE": ("FFE", "FF&E", "FURNITURE, FIXTURES & EQUIPMENT", "FF&E"),
    "DIV-15-MEP": ("DIV15", "DIVISION 15", "MEP SERVICES", "DIV-15 - MEP"),
    "OTHER": ("OTH", "OTHER", "OTHER WORKS", "OTHER"),
}
PREAMBLE: Dict[str, str] = {
    "finishes": "Supply and installation of finishes with all necessary materials including preparation, "
                "screeding / self levelling, grouting, equipment and accessories as per drawings and in "
                "accordance with the specification and Landlord requirements.",
    "joinery": "Supply and installation of joinery works complete with all necessary carcass, hardware, "
               "ironmongery, finishes and fixing accessories as per drawings and specification.",
    "doors": "Supply and installation of door works with all necessary frames, sub-frames, ironmongery and "
             "fixing accessories as per drawings and the door schedule.",
    "lighting": "Supply and installation of lighting fixtures complete with all necessary fixing accessories, "
                "fittings, supports, including conduiting and cabling all in accordance with the drawings and "
                "specification.",
    "ffe": "Supply and installation of all FF&E items complete with all necessary fixing accessories.",
    "mep": "Supply, installation, testing and commissioning of MEP services complete as per drawings and "
           "specification.",
}
UNIT_LABEL = {"m2": "m2", "lm": "lm", "nos": "Nos", "m3": "m3", "kg": "kg", "item": "Item", "ls": "LS"}
BASIS_LABEL = {"measured": "Measured", "counted": "Counted", "estimated": "Estimated", "manual": "Manual"}

COL_A, COL_B, COL_C, COL_D, COL_E, COL_F = "A", "B", "C", "D", "E", "F"
WIDTHS = {"A": 7, "B": 62, "C": 11, "D": 7.5, "E": 14, "F": 16}
HEADER_ROW = 7
FIRST_DATA_ROW = 9


@dataclass
class _TabResult:
    name: str
    total_row: int
    sections: List[BOQSection]
    summary_label: str
    line_refs: List[Tuple[str, BOQLine, BOQSection]] = field(default_factory=list)  # (ref, line, section)


def _index_letters(i: int) -> str:
    """0 -> A, 25 -> Z, 26 -> AA."""
    s = ""
    i += 1
    while i > 0:
        i, rem = divmod(i - 1, 26)
        s = chr(65 + rem) + s
    return s


def _currency(job: JobRecord) -> str:
    if job.currency:
        return job.currency
    return job.boq.summary.currency if job.boq else "INR"


# --------------------------------------------------------------------------- #
# public entry point
# --------------------------------------------------------------------------- #
def write_boq_workbook(job: JobRecord, path: Path) -> Path:
    if job.boq is None:
        raise ValueError("Job has no BOQ result to export")
    boq: BOQResult = job.boq
    ccy = _currency(job)
    date_str = datetime.now().strftime("%d-%b-%Y")

    wb = Workbook()
    wb.properties.creator = "Avighna BOQ"
    wb.properties.title = f"{job.project_name} - Bill of Quantities"
    wb.properties.subject = "Bill of Quantities"
    wb.properties.description = f"Generated by Avighna BOQ for job {job.id}"

    cover = wb.active
    cover.title = "COVER"
    _write_cover(cover, job, ccy, date_str)

    # group sections per tab, keeping division order (finishes before joinery etc.)
    by_tab: Dict[str, List[BOQSection]] = {}
    ordered = sorted(
        enumerate(boq.sections),
        key=lambda t: (DIVISION_SEQ.index(t[1].division) if t[1].division in DIVISION_SEQ else 99, t[0]),
    )
    for _, sec in ordered:
        if not sec.lines:
            continue
        by_tab.setdefault(TAB_FOR_DIVISION.get(sec.division, "OTHER"), []).append(sec)

    tabs: List[_TabResult] = []
    for tab in TAB_SEQ:
        sections = by_tab.get(tab)
        if not sections:
            continue
        tabs.append(_write_division_tab(wb, tab, sections, job, ccy))

    _write_summary(wb, job, tabs, ccy, date_str)
    _write_backup(wb, job, tabs, ccy)

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# --------------------------------------------------------------------------- #
# COVER
# --------------------------------------------------------------------------- #
def _write_cover(ws: Worksheet, job: JobRecord, ccy: str, date_str: str) -> None:
    set_widths(ws, {"A": 100})
    tab_color(ws, BRAND_RED)
    merge(ws, "A2", "AVIGHNA", font=font(30, bold=True, color=BRAND_RED), alignment=CENTER)
    put(ws, "A3", "Avighna BOQ", font=font(10, color="595959"), alignment=CENTER)
    put(ws, "A9", job.project_name.upper(), font=font(16, bold=True), alignment=CENTER)
    put(ws, "A10", (job.location or "").upper(), font=font(13, bold=True), alignment=CENTER)
    put(ws, "A13", "BILL OF QUANTITIES", font=font(24, bold=True), alignment=CENTER)
    put(ws, "A14", "INTERIOR FIT-OUT WORKS", font=font(13, bold=True), alignment=CENTER)
    put(ws, "A16", f"Level / Area: {job.level}" if job.level else None, font=font(11), alignment=CENTER)
    put(ws, "A18", f"Date: {date_str}", font=font(11), alignment=CENTER)
    vat = job.boq.summary.vat_pct if job.boq else 5.0
    put(ws, "A19", f"Currency: {ccy}     |     VAT: {vat:g}%", font=font(11), alignment=CENTER)
    put(ws, "A21", f"Prepared by: {job.created_by_email or 'Avighna BOQ'}", font=font(10), alignment=CENTER)
    put(ws, "A22", f"Approved by: {job.approved_by}" if job.approved_by else "Revision: R00",
        font=font(10), alignment=CENTER)
    put(ws, "A23", f"Drawing refs: {', '.join(job.drawing_refs)}" if job.drawing_refs else None,
        font=font(9, color="595959"), alignment=CENTER)
    put(ws, "A24", f"Job: {job.id}", font=font(9, color="595959"), alignment=CENTER)
    set_heights(ws, {2: 44, 9: 30, 10: 22, 13: 36, 14: 22})
    print_setup(ws, landscape=False, a3=False)


# --------------------------------------------------------------------------- #
# Division tab
# --------------------------------------------------------------------------- #
def _division_heading(tab: str, sections: List[BOQSection]) -> Tuple[str, str, str]:
    code, h1, h2, label = TAB_META[tab]
    divisions = {s.division for s in sections}
    if tab == "DIV-09-FINISHES" and "joinery" in divisions:
        h2 = "FINISHES & JOINERY" if "finishes" in divisions else "JOINERY"
        label = "DIV-09 - " + h2
    return h1, h2, label


def _line_text(line: BOQLine, item: Optional[FitoutItem]) -> str:
    """'J-021 - Pantry counter' + size label / specification on following lines (wrapped in the cell).

    BOQLine carries only the measured quantity; the size label and the
    brand / model specification live on the FitoutItem it was built from.
    """
    text = f"{line.code} - {line.description}" if line.code else line.description
    extras = []
    for key in ("size_label", "specification"):
        v = getattr(item, key, None) if item is not None else getattr(line, key, None)
        if v and str(v) not in text:
            extras.append(str(v))
    if extras:
        text += "\n" + "\n".join(extras)
    return text


def _project_header(ws: Worksheet, job: JobRecord, right_text: str) -> None:
    left = job.project_name.upper() + (f"\n{job.location.upper()}" if job.location else "")
    merge(ws, "A1:D2", left, font=font(10, bold=True), alignment=LEFT, border=BORDER_MEDIUM)
    merge(ws, "E1:F2", right_text, font=font(10, bold=True), alignment=CENTER, border=BORDER_MEDIUM)
    set_heights(ws, {1: 20, 2: 20})


def _column_header(ws: Worksheet, row: int, ccy: str) -> None:
    for letter, name in zip("ABCDEF", ["ITEM", "WORK DESCRIPTION", "QTY", "UNIT", "RATE", "AMOUNT"]):
        header(ws, f"{letter}{row}", name)
    for letter in "ABCD":
        put(ws, f"{letter}{row + 1}", None, font=font(9, bold=True, color=WHITE), fill=fill(BRAND_RED),
            border=BORDER_THIN, alignment=CENTER)
    for letter in "EF":
        put(ws, f"{letter}{row + 1}", f"({ccy})", font=font(9, bold=True, color=WHITE), fill=fill(BRAND_RED),
            border=BORDER_THIN, alignment=CENTER)
    set_heights(ws, {row: 20, row + 1: 16})


def _write_division_tab(wb: Workbook, tab: str, sections: List[BOQSection], job: JobRecord,
                        ccy: str) -> _TabResult:
    ws = wb.create_sheet(safe_sheet_name(tab, wb.sheetnames))
    items_by_id: Dict[str, FitoutItem] = {it.id: it for it in job.fitout_items}
    set_widths(ws, WIDTHS)
    code, _, _, _ = TAB_META[tab]
    h1, h2, label = _division_heading(tab, sections)

    _project_header(ws, job, f"BILL OF QUANTITIES\n{label}")
    merge(ws, "A4:F4", h1, font=font(18, bold=True), alignment=CENTER)
    merge(ws, "A5:F5", h2, font=font(22, bold=True, color=BRAND_RED), alignment=CENTER)
    set_heights(ws, {4: 26, 5: 32})
    _column_header(ws, HEADER_ROW, ccy)

    result = _TabResult(name=ws.title, total_row=0, sections=sections, summary_label=label)
    carried_rows: List[Tuple[str, str, int]] = []  # (page ref, section name, row)
    r = FIRST_DATA_ROW
    preambles_done = set()
    for k, sec in enumerate(sections, 1):
        page_ref = f"{code}/{k}"
        sec_start = r
        # section heading
        style_cells(ws, f"A{r}:F{r}", font=font(10, bold=True), fill=GREY_FILL, border=BORDER_THIN, alignment=LEFT)
        put(ws, f"B{r}", (sec.section or sec.division_label or sec.division).upper())
        r += 2
        # the general "supply and install ..." clause appears once per division, under its first section
        preamble = PREAMBLE.get(sec.division) if sec.division not in preambles_done else None
        preambles_done.add(sec.division)
        if preamble:
            put(ws, f"B{r}", preamble, font=font(9, italic=True), alignment=LEFT_TOP)
            ws.row_dimensions[r].height = est_row_height(preamble, WIDTHS["B"], line=12)
            r += 2
        for i, line in enumerate(sec.lines):
            ref = _index_letters(i)
            text = _line_text(line, items_by_id.get(line.item_id))
            style_cells(ws, f"A{r}:F{r}", font=font(10), border=BORDER_THIN, alignment=CENTER)
            put(ws, f"A{r}", ref, font=font(10, bold=True))
            put(ws, f"B{r}", text, alignment=LEFT_TOP)
            unit = UNIT_LABEL.get(line.unit, line.unit)
            qty_fmt = FMT_QTY_INT if line.unit in ("nos", "item", "ls") else FMT_QTY
            put(ws, f"C{r}", line.qty, number_format=qty_fmt)
            put(ws, f"D{r}", unit)
            put(ws, f"E{r}", line.rate if line.rate is not None else None, number_format=FMT_MONEY)
            put(ws, f"F{r}", f"=C{r}*E{r}", number_format=FMT_MONEY)
            ws.row_dimensions[r].height = est_row_height(text, WIDTHS["B"])
            result.line_refs.append((f"{page_ref}-{ref}", line, sec))
            r += 2
        # carried to collection
        style_cells(ws, f"A{r}:F{r}", font=font(10, bold=True), border=BORDER_THIN, alignment=CENTER, fill=PALE_FILL)
        put(ws, f"B{r}", "CARRIED TO COLLECTION", alignment=RIGHT)
        put(ws, f"C{r}", page_ref)
        put(ws, f"F{r}", f"=SUM(F{sec_start}:F{r - 1})", number_format=FMT_MONEY)
        carried_rows.append((page_ref, sec.section, r))
        r += 2

    # collection block
    style_cells(ws, f"A{r}:F{r}", font=font(10, bold=True), fill=GREY_FILL, border=BORDER_THIN, alignment=LEFT)
    put(ws, f"B{r}", "COLLECTION")
    r += 2
    coll_start = r
    for page_ref, section, crow in carried_rows:
        style_cells(ws, f"A{r}:F{r}", font=font(10), border=BORDER_THIN, alignment=CENTER)
        put(ws, f"B{r}", f"From Page {page_ref}  -  {section}", alignment=LEFT)
        put(ws, f"C{r}", page_ref)
        put(ws, f"F{r}", f"=F{crow}", number_format=FMT_MONEY)
        r += 1
    coll_end = r - 1
    r += 1
    style_cells(ws, f"A{r}:F{r}", font=font(11, bold=True), border=BORDER_MEDIUM, alignment=CENTER, fill=PALE_FILL)
    put(ws, f"B{r}", "TOTAL CARRIED TO SUMMARY", alignment=RIGHT)
    put(ws, f"C{r}", f"{code}/{len(sections) + 1}")
    put(ws, f"F{r}", f"=SUM(F{coll_start}:F{coll_end})", number_format=FMT_MONEY)
    ws.row_dimensions[r].height = 22
    result.total_row = r

    ws.freeze_panes = f"A{FIRST_DATA_ROW}"
    print_setup(ws, title_rows=f"{HEADER_ROW}:{HEADER_ROW + 1}")
    return result


# --------------------------------------------------------------------------- #
# SUMMARY
# --------------------------------------------------------------------------- #
def _write_summary(wb: Workbook, job: JobRecord, tabs: List[_TabResult], ccy: str, date_str: str) -> None:
    ws = wb.create_sheet(safe_sheet_name("SUMMARY", wb.sheetnames))
    tab_color(ws, BRAND_RED)
    set_widths(ws, {"A": 12, "B": 56, "C": 22})
    vat_pct = job.boq.summary.vat_pct if job.boq else 5.0

    left = job.project_name.upper() + (f"\n{job.location.upper()}" if job.location else "")
    merge(ws, "A1:C1", left, font=font(11, bold=True), alignment=CENTER)
    merge(ws, "A2:C2", "BILL OF QUANTITIES", font=font(11, bold=True), alignment=CENTER)
    merge(ws, "A4:C4", "GENERAL SUMMARY", font=font(14, bold=True, color=BRAND_RED), alignment=CENTER,
          border=BORDER_BOTTOM_MEDIUM)
    set_heights(ws, {1: 32, 2: 18, 4: 26})
    header(ws, "A6", "DIVISION")
    header(ws, "B6", "WORK DESCRIPTION")
    header(ws, "C6", "AMOUNT")
    for letter, val in zip("ABC", [None, None, f"({ccy})"]):
        put(ws, f"{letter}7", val, font=font(9, bold=True, color=WHITE), fill=fill(BRAND_RED),
            border=BORDER_THIN, alignment=CENTER)

    r = 9
    first = r
    for i, t in enumerate(tabs, 1):
        style_cells(ws, f"A{r}:C{r}", font=font(10), border=BORDER_THIN, alignment=CENTER)
        put(ws, f"A{r}", i)
        put(ws, f"B{r}", t.summary_label, font=font(10, bold=True), alignment=LEFT)
        put(ws, f"C{r}", f"={qs(t.name)}!F{t.total_row}", number_format=FMT_MONEY)
        ws.row_dimensions[r].height = 24
        r += 1
    last = r - 1
    r += 1

    def total_line(label: str, formula: str, *, bold_size: float = 10):
        nonlocal r
        style_cells(ws, f"A{r}:C{r}", font=font(bold_size, bold=True), border=BORDER_MEDIUM, alignment=CENTER)
        merge(ws, f"A{r}:B{r}", label, alignment=RIGHT)
        put(ws, f"C{r}", formula, number_format=FMT_MONEY)
        ws.row_dimensions[r].height = 24
        r += 1
        return r - 1

    grand = total_line("Grand Total, excluding VAT:", f"=SUM(C{first}:C{last})" if tabs else 0)
    vat = total_line(f"VAT @ {vat_pct:g}%:", f"=ROUND(C{grand}*{vat_pct:g}/100,2)")
    total_line(f"Overall Total including VAT of {vat_pct:g}%:", f"=C{grand}+C{vat}", bold_size=11)
    style_cells(ws, f"C{grand + 2}", fill=PALE_FILL)
    r += 1
    merge(ws, f"A{r}:C{r}", f"(Amount in Words)  {ccy} ........................................................................",
          font=font(10, bold=True), alignment=LEFT, border=BORDER_MEDIUM)
    ws.row_dimensions[r].height = 34
    r += 1
    merge(ws, f"A{r}:C{r}", "SIGNATURE OF TENDERER                                                            STAMP",
          font=font(10, bold=True), alignment=LEFT_BOTTOM, border=BORDER_MEDIUM)
    ws.row_dimensions[r].height = 90
    r += 2
    merge(ws, f"A{r}:C{r}", f"Generated by Avighna BOQ on {date_str}  |  Job {job.id}  |  "
                            f"Prepared by {job.created_by_email or '-'}"
                            + (f"  |  Approved by {job.approved_by}" if job.approved_by else ""),
          font=font(8, color="595959"), alignment=LEFT)
    print_setup(ws, landscape=False, a3=False)


# --------------------------------------------------------------------------- #
# MEASUREMENT BACKUP (audit trail)
# --------------------------------------------------------------------------- #
BACKUP_COLS = [
    ("Item ref", 12), ("Division", 22), ("Section", 18), ("Code", 10), ("Description", 48), ("Rooms", 26),
    ("Qty", 10), ("Unit", 7), ("Basis", 11), ("Confidence", 11), ("Source sheet", 16), ("Evidence", 55),
    ("Calculation", 40),
]


def _write_backup(wb: Workbook, job: JobRecord, tabs: List[_TabResult], ccy: str) -> None:
    ws = wb.create_sheet(safe_sheet_name("MEASUREMENT BACKUP", wb.sheetnames))
    set_widths(ws, {col(i): w for i, (_, w) in enumerate(BACKUP_COLS, 1)})
    n = len(BACKUP_COLS)
    merge(ws, f"A1:{col(n)}1", f"Measurement backup  |  {job.project_name}  |  {job.level}",
          font=font(12, bold=True), alignment=LEFT)
    put(ws, "A2", "One row per BOQ line. Basis: Measured = taken off the drawing geometry, Counted = symbols/blocks "
                  "counted, Estimated = inferred from room data, Manual = entered by the reviewer.",
        font=font(8, color="595959"), alignment=LEFT)
    for i, (name, _) in enumerate(BACKUP_COLS, 1):
        header(ws, f"{col(i)}3", name)
    ws.row_dimensions[3].height = 20

    r = 4
    for t in tabs:
        for ref, line, sec in t.line_refs:
            style_cells(ws, f"A{r}:{col(n)}{r}", font=font(9), border=BORDER_THIN, alignment=LEFT_TOP)
            desc = (f"{line.code} - " if line.code else "") + line.description
            rooms = ", ".join(line.rooms) if line.rooms else ""
            evidence = " | ".join(line.evidence) if line.evidence else ""
            values = [
                ref, sec.division_label or DIVISION_LABELS.get(sec.division, sec.division), sec.section,
                line.code or "", desc, rooms, line.qty, UNIT_LABEL.get(line.unit, line.unit),
                BASIS_LABEL.get(line.qty_basis, line.qty_basis), line.confidence, line.source_sheet or "",
                evidence, line.formula or "",
            ]
            for i, v in enumerate(values, 1):
                put(ws, f"{col(i)}{r}", v if v != "" else None)
            put(ws, f"G{r}", None, number_format=FMT_QTY_INT if line.unit in ("nos", "item", "ls") else FMT_QTY,
                alignment=CENTER)
            put(ws, f"J{r}", None, number_format=FMT_PCT, alignment=CENTER)
            ws.row_dimensions[r].height = max(
                est_row_height(desc, BACKUP_COLS[4][1]), est_row_height(evidence, BACKUP_COLS[11][1]),
                est_row_height(line.formula, BACKUP_COLS[12][1]), est_row_height(rooms, BACKUP_COLS[5][1]),
            )
            r += 1
    ws.freeze_panes = "A4"

    rooms: List[RoomRecord] = list(job.boq.rooms) if (job.boq and job.boq.rooms) else list(job.rooms)
    if rooms:
        r += 2
        merge(ws, f"A{r}:{col(n)}{r}", "Room schedule (basis for measured / estimated quantities)",
              font=font(11, bold=True), alignment=LEFT)
        r += 1
        room_cols = ["Room", "Area (m2)", "Perimeter (m)", "Ceiling ht (m)", "Floor finish", "Wall finishes",
                     "Ceiling finish", "Source sheet", "Confidence", "Evidence"]
        for i, name in enumerate(room_cols, 1):
            header(ws, f"{col(i)}{r}", name)
        r += 1
        for room in rooms:
            style_cells(ws, f"A{r}:{col(len(room_cols))}{r}", font=font(9), border=BORDER_THIN, alignment=LEFT_TOP)
            vals = [room.name, room.area_m2, room.perimeter_m, room.ceiling_height_m, room.floor_finish,
                    ", ".join(room.wall_finishes) if room.wall_finishes else None, room.ceiling_finish,
                    room.source_sheet, room.confidence, " | ".join(room.evidence) if room.evidence else None]
            for i, v in enumerate(vals, 1):
                put(ws, f"{col(i)}{r}", v)
            for letter in "BCD":
                put(ws, f"{letter}{r}", None, number_format=FMT_QTY, alignment=CENTER)
            put(ws, f"I{r}", None, number_format=FMT_PCT, alignment=CENTER)
            r += 1

    if job.boq and job.boq.warnings:
        r += 2
        put(ws, f"A{r}", "Warnings", font=font(10, bold=True))
        r += 1
        for w in job.boq.warnings:
            merge(ws, f"A{r}:{col(n)}{r}", w, font=font(9, color="C00000"), alignment=LEFT)
            r += 1
    print_setup(ws, title_rows="3:3")
