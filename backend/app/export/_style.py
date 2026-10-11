"""Shared openpyxl styling helpers for the Avighna workbooks.

Brand: header fills #ED1C24 with white bold text, body text #231F20,
Arial throughout (the reference sheets are Arial), thin borders on data.
"""
from __future__ import annotations

import math
import re
from typing import Dict, Iterable, Optional

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet

BRAND_RED = "ED1C24"
BRAND_DARK = "231F20"
WHITE = "FFFFFF"
GREY_LIGHT = "F2F2F2"
GREY_MID = "D9D9D9"
RED_PALE = "FDECEC"
FONT_NAME = "Arial"

# number formats
FMT_LEN = "0.000"          # metres, 3 decimals
FMT_KG = "0.00"            # kilograms, 2 decimals
FMT_TON = "0.000"          # tons, 3 decimals
FMT_INT = "0"
FMT_QTY = "#,##0.00"
FMT_QTY_INT = "#,##0"
FMT_MONEY = '_-* #,##0.00_-;\\-* #,##0.00_-;_-* "-"??_-;_-@_-'
FMT_PCT = "0%"


def font(size: float = 10, bold: bool = False, color: str = BRAND_DARK, italic: bool = False) -> Font:
    return Font(name=FONT_NAME, size=size, bold=bold, italic=italic, color=color)


def fill(hex_color: str) -> PatternFill:
    return PatternFill(fill_type="solid", start_color=hex_color, end_color=hex_color)


_THIN = Side(style="thin", color=BRAND_DARK)
_MEDIUM = Side(style="medium", color=BRAND_DARK)
_HAIR = Side(style="hair", color=BRAND_DARK)
BORDER_THIN = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
BORDER_MEDIUM = Border(left=_MEDIUM, right=_MEDIUM, top=_MEDIUM, bottom=_MEDIUM)
BORDER_BOTTOM = Border(bottom=_THIN)
BORDER_BOTTOM_MEDIUM = Border(bottom=_MEDIUM)
BORDER_TOP_BOTTOM = Border(top=_THIN, bottom=_MEDIUM)
NO_BORDER = Border()

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
CENTER_TOP = Alignment(horizontal="center", vertical="top", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
LEFT_TOP = Alignment(horizontal="left", vertical="top", wrap_text=True)
LEFT_BOTTOM = Alignment(horizontal="left", vertical="bottom", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center", wrap_text=False)

HEADER_FONT = font(10, bold=True, color=WHITE)
HEADER_FILL = fill(BRAND_RED)
PALE_FILL = fill(RED_PALE)
GREY_FILL = fill(GREY_LIGHT)
MID_FILL = fill(GREY_MID)


def col(n: int) -> str:
    """1 -> 'A'."""
    return get_column_letter(n)


def qs(sheet_name: str) -> str:
    """Quote a sheet name for use inside a formula: 14th Beam -> '14th Beam'."""
    return "'" + sheet_name.replace("'", "''") + "'"


_BAD_SHEET_CHARS = re.compile(r"[\[\]\:\*\?\/\\]")


def safe_sheet_name(name: str, existing: Iterable[str] = ()) -> str:
    base = _BAD_SHEET_CHARS.sub("-", name).strip() or "Sheet"
    base = base[:31]
    taken = set(existing)
    candidate = base
    i = 2
    while candidate in taken:
        suffix = f" ({i})"
        candidate = base[: 31 - len(suffix)] + suffix
        i += 1
    return candidate


def style_cells(
    ws: Worksheet,
    cell_range: str,
    *,
    font: Optional[Font] = None,
    fill: Optional[PatternFill] = None,
    border: Optional[Border] = None,
    alignment: Optional[Alignment] = None,
    number_format: Optional[str] = None,
) -> None:
    """Apply styles to every cell of a range (merged ranges need every cell bordered)."""
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for c in row:
            if font is not None:
                c.font = font
            if fill is not None:
                c.fill = fill
            if border is not None:
                c.border = border
            if alignment is not None:
                c.alignment = alignment
            if number_format is not None:
                c.number_format = number_format


def put(ws: Worksheet, ref: str, value=None, **style):
    """Write a value (None leaves the cell untouched) and style it."""
    c = ws[ref]
    if value is not None:
        c.value = value
    if style:
        style_cells(ws, ref, **style)
    return c


def merge(ws: Worksheet, cell_range: str, value=None, **style):
    """Merge a range, style every cell in it and write the value to the top-left cell."""
    if ":" in cell_range:
        ws.merge_cells(cell_range)
        top_left = cell_range.split(":")[0]
    else:
        top_left = cell_range
    if style:
        style_cells(ws, cell_range, **style)
    c = ws[top_left]
    if value is not None:
        c.value = value
    return c


def header(ws: Worksheet, cell_range: str, value=None, *, size: float = 10, number_format: Optional[str] = None):
    """Brand column header: red fill, white bold, centred, thin border."""
    return merge(
        ws, cell_range, value,
        font=font(size, bold=True, color=WHITE), fill=HEADER_FILL,
        border=BORDER_THIN, alignment=CENTER, number_format=number_format,
    )


def info(ws: Worksheet, cell_range: str, value=None, *, bold: bool = True, size: float = 12,
         alignment: Alignment = CENTER, number_format: Optional[str] = None):
    """Boxed information cell in the title block (medium border, Arial 12 bold like the reference)."""
    return merge(
        ws, cell_range, value,
        font=font(size, bold=bold), border=BORDER_MEDIUM, alignment=alignment, number_format=number_format,
    )


def set_widths(ws: Worksheet, widths: Dict[str, float]) -> None:
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width


def set_heights(ws: Worksheet, heights: Dict[int, float]) -> None:
    for row, height in heights.items():
        ws.row_dimensions[row].height = height


def tab_color(ws: Worksheet, hex_color: str) -> None:
    ws.sheet_properties.tabColor = hex_color


def print_setup(ws: Worksheet, *, landscape: bool = True, a3: bool = True,
                title_rows: Optional[str] = None) -> None:
    """Landscape A3, fit to one page wide, repeat header rows on every printed page."""
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3 if a3 else ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    if title_rows:
        ws.print_title_rows = title_rows


def est_row_height(text: Optional[str], width_chars: float, *, base: float = 15.0,
                   line: float = 12.75, min_height: float = 15.0, max_height: float = 300.0) -> float:
    """Approximate the row height needed to show wrapped text in a column of `width_chars`."""
    if not text:
        return min_height
    lines = 0
    for part in str(text).split("\n"):
        lines += max(1, int(math.ceil(len(part) / max(width_chars - 2, 1))))
    return max(min_height, min(max_height, base + line * (lines - 1)))


def blank_if_zero(v: Optional[float]):
    """The reference sheets leave zero-length legs empty rather than writing 0."""
    if v is None:
        return None
    try:
        return None if abs(float(v)) < 1e-12 else v
    except (TypeError, ValueError):
        return v
