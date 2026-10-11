"""Excel export of the fit-out Bill of Quantities (KR-style layout)."""
from __future__ import annotations

from typing import Dict, List, Tuple

import openpyxl
import pytest

from app.export import write_workbook
from app.export.boq_xlsx import write_boq_workbook
from app.models.schemas import (
    DIVISION_LABELS, BOQLine, BOQResult, BOQSection, BOQSummary, FitoutItem, JobRecord, Module, RoomRecord,
)


def make_items() -> List[FitoutItem]:
    return [
        FitoutItem(division="finishes", section="Floor Finishes", code="FC-201", description="Cream Cremo Delicato stone flooring",
                   unit="m2", qty=270, qty_basis="measured", rooms=["Reception", "Lounge"],
                   specification="Honed finish, 20 mm thick, laid on screed", confidence=0.85,
                   evidence=["FC-201 hatch area 270.4 m2"], source_sheet="ID-201 Floor Finishes Plan"),
        FitoutItem(division="finishes", section="Floor Finishes", code="FC-202", description="600 x 600 mm carpet tiles",
                   unit="m2", qty=175, qty_basis="measured", rooms=["Open Office"], confidence=0.8,
                   evidence=["FC-202 hatch 175 m2"], source_sheet="ID-201 Floor Finishes Plan"),
        FitoutItem(division="finishes", section="Wall Finishes", code="WC-103", description="Veneer wall panelling",
                   unit="m2", qty=88, qty_basis="estimated", rooms=["Reception"], confidence=0.6,
                   evidence=["WC-103 tag on wall elevation"], source_sheet="ID-301 Wall Finishes Plan"),
        FitoutItem(division="finishes", section="Ceilings", code="CL-301", description="Paint finish to gypsum ceiling",
                   unit="m2", qty=308, qty_basis="estimated", rooms=["Reception", "Lounge", "Open Office"],
                   confidence=0.7, evidence=["CL-301 ceiling legend"], source_sheet="ID-401 RCP"),
        FitoutItem(division="joinery", section="Joinery", code="J-021", description="Pantry counter",
                   unit="nos", qty=1, qty_basis="counted", rooms=["Pantry"], size_label="3225mm L x 650mm W x 2700mm H",
                   specification="ST-405 reconstituted stone top, VN-701 veneer carcass", confidence=0.9,
                   evidence=["J-021 block on furniture plan"], source_sheet="ID-501 Furniture Plan"),
        FitoutItem(division="joinery", section="Joinery", code="J-022", description="Banquet seating",
                   unit="nos", qty=1, qty_basis="counted", rooms=["Pantry"], confidence=0.9,
                   evidence=["J-022 block"], source_sheet="ID-501 Furniture Plan"),
        FitoutItem(division="lighting", section="Lighting", code="LI-501", description="Pendant light",
                   unit="nos", qty=3, qty_basis="counted", rooms=["Reception"],
                   specification="A-N-D Column Pendant 175 Horizontal, 2700K", confidence=0.95,
                   evidence=["LI-501 x3 on RCP"], source_sheet="ID-401 RCP"),
        FitoutItem(division="ffe", section="Chairs", code="CH-101", description="Reception armchair",
                   unit="nos", qty=7, qty_basis="counted", rooms=["Reception"], rate=1200.0, confidence=0.9,
                   evidence=["CH-101 x7 blocks"], source_sheet="ID-501 Furniture Plan"),
        FitoutItem(division="ffe", section="Tables", code="TB-201", description="Meeting table",
                   unit="nos", qty=1, qty_basis="counted", rooms=["Meeting Room"], confidence=0.9,
                   evidence=["TB-201 block"], source_sheet="ID-501 Furniture Plan"),
    ]


def make_rooms() -> List[RoomRecord]:
    return [
        RoomRecord(name="Reception", area_m2=120.5, perimeter_m=46.0, ceiling_height_m=3.0, floor_finish="FC-201",
                   wall_finishes=["WC-103"], ceiling_finish="CL-301", source_sheet="ID-101 GA Plan", confidence=0.8),
        RoomRecord(name="Lounge", area_m2=150.0, perimeter_m=52.0, ceiling_height_m=3.0, floor_finish="FC-201",
                   ceiling_finish="CL-301", source_sheet="ID-101 GA Plan", confidence=0.8),
        RoomRecord(name="Open Office", area_m2=175.0, perimeter_m=60.0, ceiling_height_m=2.8, floor_finish="FC-202",
                   ceiling_finish="CL-301", source_sheet="ID-101 GA Plan", confidence=0.7),
    ]


def build_boq(items: List[FitoutItem], rooms: List[RoomRecord]) -> BOQResult:
    groups: Dict[Tuple[str, str], List[FitoutItem]] = {}
    for it in items:
        groups.setdefault((it.division, it.section), []).append(it)
    sections: List[BOQSection] = []
    seq = 0
    for (division, section), its in groups.items():
        lines = []
        for i, it in enumerate(its):
            seq += 1
            amount = it.qty * it.rate if it.rate is not None else None
            lines.append(BOQLine(
                seq=seq, item_ref=chr(65 + i), item_id=it.id, division=division, section=section, code=it.code,
                description=it.description, unit=it.unit, qty=it.qty, rate=it.rate, amount=amount,
                qty_basis=it.qty_basis, rooms=it.rooms, confidence=it.confidence, evidence=it.evidence,
                source_sheet=it.source_sheet, formula=f"{it.qty_basis}: {it.qty:g} {it.unit}",
            ))
        qty_by_unit: Dict[str, float] = {}
        for ln in lines:
            qty_by_unit[ln.unit] = qty_by_unit.get(ln.unit, 0) + ln.qty
        sections.append(BOQSection(division=division, division_label=DIVISION_LABELS[division], section=section,
                                   lines=lines, qty_by_unit=qty_by_unit,
                                   amount=sum(l.amount or 0 for l in lines)))
    subtotal = sum(s.amount for s in sections)
    summary = BOQSummary(currency="AED", vat_pct=5.0, line_count=len(items), priced_lines=1,
                         subtotal=subtotal, vat=subtotal * 0.05, grand_total=subtotal * 1.05)
    return BOQResult(sections=sections, rooms=rooms, summary=summary, warnings=["Ceiling area estimated from room areas"])


@pytest.fixture(scope="module")
def job() -> JobRecord:
    items, rooms = make_items(), make_rooms()
    return JobRecord(
        id="fit0001", module=Module.fitout, project_name="Proposed Avighna Group Office and Sales Center",
        level="Level 12", currency="AED", location="Dubai, U.A.E.", drawing_refs=["ID-101", "ID-201", "ID-401", "ID-501"],
        fitout_items=items, rooms=rooms, boq=build_boq(items, rooms), created_by_email="qs@avighna.in",
    )


@pytest.fixture(scope="module")
def workbook(job, tmp_path_factory):
    out = tmp_path_factory.mktemp("boq") / "boq.xlsx"
    assert write_boq_workbook(job, out) == out
    return openpyxl.load_workbook(out)


def find_row(ws, col_letter: str, predicate, start: int = 1) -> int:
    for r in range(start, ws.max_row + 1):
        v = ws[f"{col_letter}{r}"].value
        if v is not None and predicate(v):
            return r
    raise AssertionError(f"no match in column {col_letter} of {ws.title}")


def test_sheet_names(workbook):
    assert workbook.sheetnames == ["COVER", "DIV-09-FINISHES", "DIV-13-LIGHTING", "FFE", "SUMMARY", "MEASUREMENT BACKUP"]
    assert workbook.properties.creator == "Avighna BOQ"


def test_cover(workbook):
    ws = workbook["COVER"]
    values = [ws[f"A{r}"].value for r in range(1, 26)]
    assert "AVIGHNA" in values
    assert "PROPOSED AVIGHNA GROUP OFFICE AND SALES CENTER" in values
    assert "DUBAI, U.A.E." in values
    assert "BILL OF QUANTITIES" in values
    assert any(isinstance(v, str) and v.startswith("Date:") for v in values)


def test_finishes_tab_layout(workbook, job):
    ws = workbook["DIV-09-FINISHES"]
    assert ws["A1"].value == "PROPOSED AVIGHNA GROUP OFFICE AND SALES CENTER\nDUBAI, U.A.E."
    assert ws["E1"].value == "BILL OF QUANTITIES\nDIV-09 - FINISHES & JOINERY"
    assert ws["A4"].value == "DIVISION 09" and ws["A5"].value == "FINISHES & JOINERY"
    assert [ws.cell(7, c).value for c in range(1, 7)] == ["ITEM", "WORK DESCRIPTION", "QTY", "UNIT", "RATE", "AMOUNT"]
    assert ws["E8"].value == "(AED)" and ws["F8"].value == "(AED)"
    assert ws["A7"].fill.fgColor.rgb.endswith("ED1C24")
    assert ws["A7"].font.bold and ws["A7"].font.color.rgb.endswith("FFFFFF")
    assert ws.freeze_panes == "A9"
    assert ws.page_setup.orientation == "landscape" and int(ws.page_setup.paperSize) == int(ws.PAPERSIZE_A3)

    # section headings in order: finishes first, joinery after
    headings = [ws[f"B{r}"].value for r in range(9, ws.max_row + 1)
                if ws[f"B{r}"].value in ("FLOOR FINISHES", "WALL FINISHES", "CEILINGS", "JOINERY")]
    assert headings == ["FLOOR FINISHES", "WALL FINISHES", "CEILINGS", "JOINERY"]

    # line rows: ref letters restart per section, description = code - description (+ spec), amount formula
    r = find_row(ws, "B", lambda v: isinstance(v, str) and v.startswith("FC-201 - "))
    assert ws[f"A{r}"].value == "A"
    assert ws[f"B{r}"].value == "FC-201 - Cream Cremo Delicato stone flooring\nHoned finish, 20 mm thick, laid on screed"
    assert ws[f"C{r}"].value == 270 and ws[f"D{r}"].value == "m2"
    assert ws[f"E{r}"].value is None
    assert ws[f"F{r}"].value == f"=C{r}*E{r}"
    assert ws[f"B{r}"].alignment.wrap_text
    r2 = find_row(ws, "B", lambda v: isinstance(v, str) and v.startswith("FC-202 - "))
    assert ws[f"A{r2}"].value == "B"
    r3 = find_row(ws, "B", lambda v: isinstance(v, str) and v.startswith("WC-103 - "))
    assert ws[f"A{r3}"].value == "A"
    rj = find_row(ws, "B", lambda v: isinstance(v, str) and v.startswith("J-021 - "))
    assert ws[f"A{rj}"].value == "A"
    assert "3225mm L x 650mm W x 2700mm H" in ws[f"B{rj}"].value
    assert ws[f"D{rj}"].value == "Nos"

    # carried to collection per section with SUM formulas and page refs
    carried = [r for r in range(9, ws.max_row + 1) if ws[f"B{r}"].value == "CARRIED TO COLLECTION"]
    assert len(carried) == 4
    assert [ws[f"C{r}"].value for r in carried] == ["DIV09/1", "DIV09/2", "DIV09/3", "DIV09/4"]
    for r in carried:
        assert ws[f"F{r}"].value.startswith("=SUM(F") and ws[f"F{r}"].value.endswith(f":F{r - 1})")
    # collection block and tab total
    coll = find_row(ws, "B", lambda v: v == "COLLECTION")
    first_from = coll + 2
    assert ws[f"B{first_from}"].value.startswith("From Page DIV09/1")
    assert ws[f"F{first_from}"].value == f"=F{carried[0]}"
    total = find_row(ws, "B", lambda v: v == "TOTAL CARRIED TO SUMMARY")
    assert ws[f"F{total}"].value == f"=SUM(F{first_from}:F{first_from + 3})"


def test_rate_written_when_present(workbook):
    ws = workbook["FFE"]
    assert ws["A4"].value == "FF&E"
    r = find_row(ws, "B", lambda v: isinstance(v, str) and v.startswith("CH-101 - "))
    assert ws[f"C{r}"].value == 7 and ws[f"E{r}"].value == 1200.0 and ws[f"F{r}"].value == f"=C{r}*E{r}"
    headings = [ws[f"B{r}"].value for r in range(9, ws.max_row + 1) if ws[f"B{r}"].value in ("CHAIRS", "TABLES")]
    assert headings == ["CHAIRS", "TABLES"]
    lw = workbook["DIV-13-LIGHTING"]
    assert lw["A4"].value == "DIVISION 13" and lw["A5"].value == "LIGHTING"
    r = find_row(lw, "B", lambda v: isinstance(v, str) and v.startswith("LI-501 - "))
    assert "A-N-D Column Pendant 175 Horizontal" in lw[f"B{r}"].value
    assert lw[f"C{r}"].value == 3


def test_summary_tab(workbook):
    ws = workbook["SUMMARY"]
    assert ws["A4"].value == "GENERAL SUMMARY"
    assert (ws["A6"].value, ws["B6"].value, ws["C6"].value) == ("DIVISION", "WORK DESCRIPTION", "AMOUNT")
    assert ws["C7"].value == "(AED)"
    labels = [ws[f"B{r}"].value for r in (9, 10, 11)]
    assert labels == ["DIV-09 - FINISHES & JOINERY", "DIV-13 - LIGHTING", "FF&E"]
    fin = workbook["DIV-09-FINISHES"]
    fin_total = find_row(fin, "B", lambda v: v == "TOTAL CARRIED TO SUMMARY")
    assert ws["C9"].value == f"='DIV-09-FINISHES'!F{fin_total}"
    assert ws["C10"].value.startswith("='DIV-13-LIGHTING'!F")
    assert ws["C11"].value.startswith("=FFE!F") or ws["C11"].value.startswith("='FFE'!F")
    grand = find_row(ws, "A", lambda v: v == "Grand Total, excluding VAT:")
    assert ws[f"C{grand}"].value == "=SUM(C9:C11)"
    assert ws[f"A{grand + 1}"].value == "VAT @ 5%:"
    assert ws[f"C{grand + 1}"].value == f"=ROUND(C{grand}*5/100,2)"
    assert ws[f"A{grand + 2}"].value == "Overall Total including VAT of 5%:"
    assert ws[f"C{grand + 2}"].value == f"=C{grand}+C{grand + 1}"
    assert ws[f"A{grand + 4}"].value.startswith("(Amount in Words)  AED")
    assert ws[f"A{grand + 5}"].value.startswith("SIGNATURE OF TENDERER")
    assert ws.sheet_properties.tabColor.rgb.endswith("ED1C24")


def test_measurement_backup(workbook, job):
    ws = workbook["MEASUREMENT BACKUP"]
    assert [ws.cell(3, c).value for c in range(1, 14)] == [
        "Item ref", "Division", "Section", "Code", "Description", "Rooms", "Qty", "Unit", "Basis", "Confidence",
        "Source sheet", "Evidence", "Calculation",
    ]
    rows = []
    r = 4
    while ws[f"A{r}"].value:  # contiguous block of BOQ lines; the room schedule follows after a gap
        rows.append(r)
        r += 1
    assert len(rows) == 9
    first = rows[0]
    assert ws[f"A{first}"].value == "DIV09/1-A"
    assert ws[f"B{first}"].value == "DIV-09 - FINISHES" and ws[f"C{first}"].value == "Floor Finishes"
    assert ws[f"D{first}"].value == "FC-201"
    assert ws[f"F{first}"].value == "Reception, Lounge"
    assert ws[f"G{first}"].value == 270 and ws[f"H{first}"].value == "m2" and ws[f"I{first}"].value == "Measured"
    assert ws[f"J{first}"].value == 0.85 and ws[f"J{first}"].number_format == "0%"
    assert ws[f"K{first}"].value == "ID-201 Floor Finishes Plan"
    assert ws[f"L{first}"].value == "FC-201 hatch area 270.4 m2"
    # joinery lines carry the DIV09 page ref of their own section
    j = next(r for r in rows if ws[f"D{r}"].value == "J-021")
    assert ws[f"A{j}"].value == "DIV09/4-A" and ws[f"B{j}"].value == "DIV-09 - JOINERY"
    # room schedule follows
    room_hdr = find_row(ws, "A", lambda v: v == "Room")
    assert [ws[f"A{room_hdr + i}"].value for i in (1, 2, 3)] == ["Reception", "Lounge", "Open Office"]
    assert ws[f"B{room_hdr + 1}"].value == 120.5
    warn = find_row(ws, "A", lambda v: v == "Warnings")
    assert ws[f"A{warn + 1}"].value == "Ceiling area estimated from room areas"


def test_dispatch_on_module(job, tmp_path):
    out = write_workbook(job, tmp_path / "d.xlsx")
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames[0] == "COVER"
    openpyxl.load_workbook(out, data_only=True)


def test_preliminaries_and_doors_tabs(job, tmp_path):
    items = make_items() + [
        FitoutItem(division="preliminaries", section="General Requirements", description="Site mobilisation and hoarding",
                   unit="ls", qty=1, qty_basis="manual"),
        FitoutItem(division="doors", section="Door Works", code="D-01", description="Double leaf pivot glass door",
                   unit="nos", qty=1, qty_basis="counted", size_label="2000mm L x 2700mm H"),
    ]
    j2 = job.model_copy(update={"boq": build_boq(items, make_rooms())})
    wb = openpyxl.load_workbook(write_boq_workbook(j2, tmp_path / "p.xlsx"))
    assert wb.sheetnames == ["COVER", "DIV-01-PRELIMINARIES", "DIV-09-FINISHES", "DIV-08-DOORS", "DIV-13-LIGHTING",
                             "FFE", "SUMMARY", "MEASUREMENT BACKUP"]
    pre = wb["DIV-01-PRELIMINARIES"]
    assert pre["A5"].value == "PRELIMINARIES & GENERAL REQUIREMENTS"
    r = find_row(pre, "B", lambda v: v == "Site mobilisation and hoarding")
    assert pre[f"D{r}"].value == "LS" and pre[f"C{r}"].value == 1
    s = wb["SUMMARY"]
    assert [s[f"B{r}"].value for r in range(9, 14)] == [
        "DIV-01 - PRELIMINARIES", "DIV-09 - FINISHES & JOINERY", "DIV-08 - DOORS", "DIV-13 - LIGHTING", "FF&E",
    ]
