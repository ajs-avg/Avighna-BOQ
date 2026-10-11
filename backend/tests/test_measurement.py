"""Measurement engine vs the Hinglish guide examples; Excel export; AI-path through the API with a fake client."""
from __future__ import annotations

import time

import fitz
import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.ai import gemini as gm
from app.ai.prompts import (
    AIMAssumptions, AIMeasurementExtraction, AIMMisc, AIMOpening, AIMRoom, AISheetClass, AIMWall,
)
from app.engine.measurement import compute_measurement, material_totals
from app.export import write_workbook
from app.main import app
from app.models.schemas import (
    JobRecord, MeasurementExtraction, MemberEl, MiscEl, Module, OpeningEl, ProjectAssumptions, RoomEl, SlabEl, WallEl,
)


def _guide_bedroom() -> MeasurementExtraction:
    return MeasurementExtraction(
        assumptions=ProjectAssumptions(floor_height_m=3.65, wall_thickness_m=0.2, construction="mivan", concrete_grade="M40"),
        rooms=[RoomEl(name="Bedroom 1", level="6th floor", length_m=4.0, width_m=3.0, height_m=3.0, plastered=True,
                      floor_finish="vitrified tile flooring 600x600 mm")],
        openings=[OpeningEl(tag="D1", kind="door", width_m=1.0, height_m=2.1, room="Bedroom 1", wall_tag="W2"),
                  OpeningEl(tag="W1", kind="window", width_m=1.8, height_m=1.2, room="Bedroom 1", wall_tag="W1")],
        walls=[WallEl(tag=t, wall_type="rcc", length_m=l, thickness_m=0.2, height_m=3.65, room="Bedroom 1", level="6th floor")
               for t, l in (("W1", 4.9), ("W3", 4.9), ("W2", 3.6), ("W4", 3.6))],
    )


def _block(res, code):
    return next(b for b in res.blocks if b.item_code == code)


def test_guide_example_1a_walls_and_formwork():
    res = compute_measurement(_guide_bedroom())
    assert _block(res, "RCC-WALL").total == 11.558          # guide: 11.56 cum (3.577+3.577+2.628+2.628-0.42-0.432)
    assert _block(res, "FW-WALL").total == 115.58            # guide: 115.58 sqm
    assert _block(res, "DOOR-D1").total == 1
    steel = _block(res, "STEEL")
    assert steel.total == round(11.558 * 125, 3) and steel.rows[0].rule.startswith("thumb rule")
    rcc_first = [b.item_code for b in res.blocks][:3]
    assert rcc_first == ["RCC-WALL", "FW-WALL", "STEEL"]      # concrete before formwork, steel last in RCC


def test_guide_example_a_room_finishing():
    """Chapter 9: 4.00 x 3.00 x 3.00 room, door 1.00x2.10, window 1.50x1.20 -> 12 sqm floor, 13 rmt skirting, 38.10 sqm plaster."""
    ex = MeasurementExtraction(
        rooms=[RoomEl(name="Bedroom 1", length_m=4.0, width_m=3.0, height_m=3.0, plastered=True)],
        openings=[OpeningEl(tag="D1", kind="door", width_m=1.0, height_m=2.1, room="Bedroom 1"),
                  OpeningEl(tag="W1", kind="window", width_m=1.5, height_m=1.2, room="Bedroom 1")])
    res = compute_measurement(ex)
    floor = next(b for b in res.blocks if b.item_code.startswith("FLOOR"))
    assert floor.total == 12.0
    assert _block(res, "SKIRT").total == 13.0
    assert _block(res, "PLASTER-INT").total == 38.1
    assert _block(res, "PLASTER-CLG").total == 12.0
    assert _block(res, "PAINT-WALL").total == 38.1 and _block(res, "PAINT-CLG").total == 12.0
    assert _block(res, "WIN-W1").total == 1.8
    rules = [r.rule for r in _block(res, "PLASTER-INT").rows]
    assert any("one face" in x for x in rules)


def test_big_opening_adds_reveals_and_wet_room_items():
    ex = MeasurementExtraction(
        assumptions=ProjectAssumptions(wall_thickness_m=0.23),
        rooms=[RoomEl(name="Living", length_m=6.0, width_m=4.0, height_m=3.0, plastered=True),
               RoomEl(name="Toilet 1", length_m=2.0, width_m=1.5, height_m=3.0, plastered=True, wet_area=True, dado_height_m=2.1,
                      wall_finish="ceramic wall tiles")],
        openings=[OpeningEl(tag="SW1", kind="window", width_m=3.0, height_m=2.4, room="Living")])   # 7.2 sqm > 3
    res = compute_measurement(ex)
    pl = _block(res, "PLASTER-INT")
    living = [r for r in pl.rows if "Living" in r.location or "SW1" in r.location]
    assert any(r.qty < 0 and r.nos == -2 for r in living)            # both faces deducted
    assert any(r.location.startswith("Add: jambs") for r in living)  # reveals added
    assert _block(res, "DADO").total == round(2 * (2.0 + 1.5) * 2.1, 3)
    assert _block(res, "WP").total == round(3.0 + 2 * (2.0 + 1.5) * 0.3, 3)
    assert not any("Toilet 1" in r.location for r in _block(res, "PAINT-WALL").rows)   # tiled walls not painted


def test_members_slabs_and_materials():
    ex = MeasurementExtraction(
        slabs=[SlabEl(tag="S1", length_m=5.0, width_m=4.0, thickness_m=0.15, beam_width_m=0.23)],
        members=[MemberEl(tag="C1", kind="column", width_m=0.3, depth_m=0.6, length_m=3.5),
                 MemberEl(tag="B1", kind="beam", width_m=0.23, depth_m=0.45, length_m=4.7, slab_thickness_m=0.15)],
        misc=[MiscEl(description="Wash basin with CP fittings", unit="nos", qty=2, trade="Sanitary")])
    res = compute_measurement(ex)
    assert _block(res, "RCC-COL").total == 0.63 and _block(res, "FW-COL").total == 6.3        # guide ch.10
    assert _block(res, "RCC-BEAM").total == 0.324 and _block(res, "FW-BEAM").total == 3.901  # 0.3243 cum, 3.90 sqm
    assert _block(res, "RCC-SLAB").total == 3.0 and _block(res, "FW-SLAB").total == round(4.77 * 3.77, 3)
    mats = material_totals(res.materials)
    assert "Ready-mix concrete" in mats and "TMT bars Fe500D" in mats
    assert any(l.description.startswith("Wash basin") and l.unit == "nos" and l.qty == 2 for l in res.boq)


def test_measurement_workbook(tmp_path):
    ex = _guide_bedroom()
    job = JobRecord(id="m1", project_name="Sample", level="6th floor", module=Module.measurement, elements=ex)
    job.measurement = compute_measurement(ex)
    out = write_workbook(job, tmp_path / "m.xlsx")
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Measurement Sheet", "BOQ (Abstract)", "Material Requirement", "Tender vs Site", "Elements (from drawing)"]
    ws = wb["Measurement Sheet"]
    formulas = [c.value for row in ws.iter_rows(min_col=7, max_col=7) for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert any("*" in f for f in formulas) and any(f.startswith("=SUM(") for f in formulas)
    bq = wb["BOQ (Abstract)"]
    links = [c.value for row in bq.iter_rows(min_col=3, max_col=3) for c in row if isinstance(c.value, str) and c.value.startswith("='Measurement Sheet'")]
    assert len(links) == len(job.measurement.boq)


# ---------------------------------------------------------------- AI path through the API (fake Gemini)
class FakeGemini:
    enabled, model_fast, model_pro, message = True, "fake-flash", "fake-pro", "fake"

    def status(self):
        from app.models.schemas import AIStatus
        return AIStatus(enabled=True, provider="gemini", model_fast="fake-flash", model_pro="fake-pro", message="fake")

    def generate_json(self, schema, prompt, images=(), tier="pro", **kw):
        if schema is AISheetClass:
            return AISheetClass(sheet_type="architectural", discipline="architectural", units="mm", relevant_for_boq=True, confidence=0.9)
        if schema is AIMeasurementExtraction:
            return AIMeasurementExtraction(
                rooms=[AIMRoom(name="Bedroom 1", level="6th floor", length_mm=4000, width_mm=3000, height_mm=3000, floor_finish="vitrified tiles 600x600",
                              confidence=0.9, evidence=["BEDROOM 4000 x 3000"])],
                openings=[AIMOpening(tag="D1", kind="door", width_mm=1000, height_mm=2100, room="Bedroom 1", description="flush door", confidence=0.9),
                          AIMOpening(tag="W1", kind="window", width_mm=1500, height_mm=1200, room="Bedroom 1", confidence=0.8)],
                walls=[AIMWall(tag="W1", wall_type="block", length_mm=4000, thickness_mm=150, height_mm=3000, room="Bedroom 1")],
                misc=[AIMMisc(description="Ceiling fan point", unit="nos", qty=1, trade="Electrical", room="Bedroom 1")],
                assumptions=AIMAssumptions(floor_height_mm=3000, wall_thickness_mm=150, construction="conventional", concrete_grade="M30"))
        raise AssertionError(schema)


@pytest.fixture(scope="module")
def client():
    from _pytest.monkeypatch import MonkeyPatch
    mp = MonkeyPatch()
    mp.setattr(gm.GeminiClient, "get", classmethod(lambda cls: FakeGemini()))
    with TestClient(app) as c:
        yield c
    mp.undo()


def _token(client):
    for pw in ("Test-Admin-123", "New-Admin-Pass-1"):
        r = client.post("/api/v1/auth/login", json={"email": "admin@test.local", "password": pw})
        if r.status_code == 200:
            return {"Authorization": f"Bearer {r.json()['access_token']}"}
    raise AssertionError("login failed")


def test_measurement_job_end_to_end(client, tmp_path):
    h = _token(client)
    doc = fitz.open(); page = doc.new_page(width=1190, height=842); page.insert_text((60, 60), "FLOOR PLAN 6TH FLOOR  BEDROOM 4000 x 3000", fontsize=18)
    p = tmp_path / "plan.pdf"; doc.save(str(p)); doc.close()
    with open(p, "rb") as fh:
        r = client.post("/api/v1/jobs", data={"project_name": "Sample Tower", "level": "6th floor"},
                        files=[("files", ("plan.pdf", fh, "application/pdf"))], headers=h)
    assert r.status_code == 202, r.text
    jid = r.json()["job_id"]
    for _ in range(120):
        j = client.get(f"/api/v1/jobs/{jid}", headers=h).json()
        if j["status"] in ("ready", "failed"):
            break
        time.sleep(0.5)
    assert j["status"] == "ready", j.get("error")
    assert j["module"] == "measurement" and j["elements"]["rooms"][0]["name"] == "Bedroom 1"
    codes = {b["item_code"]: b["total"] for b in j["measurement"]["blocks"]}
    assert codes["SKIRT"] == 13.0 and codes["PLASTER-INT"] == 38.1 and codes["DOOR-D1"] == 1
    assert any(l["trade"] == "Electrical" for l in j["measurement"]["boq"])
    # reviewer edits a room size -> recomputed
    els = j["elements"]; els["rooms"][0]["length_m"] = 5.0
    r = client.put(f"/api/v1/jobs/{jid}/elements", json={"elements": els}, headers=h)
    assert r.status_code == 200
    codes2 = {b["item_code"]: b["total"] for b in r.json()["measurement"]["blocks"]}
    assert codes2["SKIRT"] == 15.0
    r = client.get(f"/api/v1/jobs/{jid}/export.xlsx", headers=h)
    assert r.status_code == 200 and len(r.content) > 8000
    lst = client.get("/api/v1/jobs?module=measurement", headers=h).json()
    assert lst["items"][0]["headline"].endswith("sqm paint") or "BOQ items" in lst["items"][0]["headline"]


def test_openings_match_rooms_by_schedule_text() -> None:
    """A door schedule line like 'BOARDROOM, BOOTH, OFFICES x6' must deduct in every matching room."""
    from app.engine.measurement import compute_measurement
    from app.models.schemas import MeasurementExtraction, OpeningEl, RoomEl

    rooms = [RoomEl(id="r1", name="BOARDROOM", length_m=5.49, width_m=6.26, height_m=2.85),
             RoomEl(id="r2", name="BOOTH 2", length_m=2.0, width_m=1.39, height_m=2.85),
             RoomEl(id="r3", name="SALES OFFICE", length_m=8.8, width_m=10.7, height_m=2.85)]
    ops = [OpeningEl(id="o1", tag="D06", kind="door", width_m=1.0, height_m=2.7, count=6, room="BOARDROOM, BOOTH, OFFICES"),
           OpeningEl(id="o2", tag="D05", kind="door", width_m=1.8, height_m=2.7, count=1, room="RECEPTION")]
    res = compute_measurement(MeasurementExtraction(rooms=rooms, openings=ops))
    skirt = next(b for b in res.blocks if b.item_code == "SKIRT")
    ded = [r for r in skirt.rows if r.deduction]
    assert len(ded) == 3 and sorted(r.nos for r in ded) == [-2.0, -2.0, -2.0]
    assert any("D05" in w and "RECEPTION" in w for w in res.warnings)
