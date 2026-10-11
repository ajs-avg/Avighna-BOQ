"""Excel export of the Bar Bending Schedule: layout, formulas and arithmetic re-checked from the cells."""
from __future__ import annotations

from pathlib import Path
from typing import Dict, Optional

import openpyxl
import pytest
from openpyxl.utils import get_column_letter

from app.engine import compute_bbs
from app.export import write_workbook
from app.export.bbs_xlsx import write_bbs_workbook
from app.models.schemas import (
    BarGroup, BeamSpec, ColumnSpec, ColumnTie, JobRecord, Module, SlabLayer, SlabSpec, StairFlight,
    StairSpec, Standards, StirrupZone, WallSpec,
)

LEVEL = "14th Floor"


def make_members():
    return [
        BeamSpec(
            tag="B1A", nos=2, span_m=2.45, width_m=0.23, depth_m=0.75,
            bottom_bars=[BarGroup(count=2, dia_mm=20)], top_bars=[BarGroup(count=2, dia_mm=20)],
            extra_top_left=[BarGroup(count=2, dia_mm=16)], extra_top_right=[BarGroup(count=2, dia_mm=16)],
            stirrups=[StirrupZone(dia_mm=8, spacing_mm=100, zone="end"),
                      StirrupZone(dia_mm=8, spacing_mm=150, zone="mid")],
            source_sheet="9680/P.P/13", confidence=0.9, evidence=["B1A 230x750", "2-T20 BOT / 2-T20 TOP"],
        ),
        BeamSpec(
            tag="B1B", nos=1, span_m=2.705, width_m=0.3, depth_m=0.425,
            bottom_bars=[BarGroup(count=3, dia_mm=12)], bottom_curtail=[BarGroup(count=1, dia_mm=12)],
            top_bars=[BarGroup(count=3, dia_mm=12)], extra_top_left=[BarGroup(count=2, dia_mm=16)],
            stirrups=[StirrupZone(dia_mm=8, spacing_mm=125, zone="full")],
            left_support="end", right_support="continuous",
            source_sheet="9680/P.P/13", confidence=0.8, evidence=["B1B 300x425"],
        ),
        SlabSpec(
            tag="S3 Near C6", nos=1, length_m=4.9, width_m=5.35, thickness_m=0.175,
            short_main=SlabLayer(dia_mm=8, spacing_mm=150, bent_up=True),
            long_main=SlabLayer(dia_mm=8, spacing_mm=175, bent_up=True),
            short_top_dist_end=SlabLayer(dia_mm=8, spacing_mm=150),
            long_top_dist_end=SlabLayer(dia_mm=8, spacing_mm=175),
            source_sheet="9680/P.P/13A", confidence=0.7, evidence=["S3 175 THK"],
        ),
        ColumnSpec(
            tag="C6", nos=2, length_m=0.3, width_m=0.6, height_m=3.8,
            vertical_bars=[BarGroup(count=8, dia_mm=16)],
            ties=[ColumnTie(dia_mm=8, spacing_mm=100, kind="main_ring"),
                  ColumnTie(dia_mm=8, spacing_mm=100, kind="inner_tie", per_layer=2)],
            source_sheet="9680/P.P/6", confidence=0.85, evidence=["C6 300x600 8-T16"],
        ),
        WallSpec(
            tag="NSW1 On B36", nos=2, length_m=1.5, thickness_m=0.15, height_m=3.8,
            vertical=SlabLayer(dia_mm=8, spacing_mm=300), horizontal=SlabLayer(dia_mm=8, spacing_mm=200), faces=2,
            source_sheet="9680/P.P/13D", confidence=0.6, evidence=["150 THK PARDI"],
        ),
        StairSpec(
            tag="ST1", nos=1,
            flights=[StairFlight(name="Flight 1", width_m=1.575, waist_m=0.175, going_m=2.7, rise_m=1.5,
                                 landing_a_m=1.2, landing_b_m=1.15,
                                 main=SlabLayer(dia_mm=12, spacing_mm=100), dist=SlabLayer(dia_mm=8, spacing_mm=150))],
            source_sheet="9680/P.P/11B", confidence=0.75, evidence=["STAIR 1 WAIST 175"],
        ),
    ]


@pytest.fixture(scope="module")
def job() -> JobRecord:
    std = Standards()
    members = make_members()
    result = compute_bbs(members, std, LEVEL)
    return JobRecord(
        id="job1234abcd", module=Module.structural, project_name="Nishika Elysium (Avighna)", level=LEVEL,
        location="Worli, Mumbai", contractor="Mehta Jaising Builders LLP", drawn_by="MUSHARRAF",
        drawing_refs=["9680/P.P/13", "13D", "13A", "13B", "201"], standards=std, members=members, result=result,
        created_by_email="engineer@avighna.in",
    )


@pytest.fixture(scope="module")
def workbook(job, tmp_path_factory):
    out = tmp_path_factory.mktemp("bbs") / "bbs.xlsx"
    written = write_bbs_workbook(job, out)
    assert written == out and out.exists() and out.stat().st_size > 0
    wb = openpyxl.load_workbook(out)  # formulas as strings
    return wb


# --------------------------------------------------------------------------- helpers
def find_row(ws, col_letter: str, text: str, start: int = 1) -> int:
    for r in range(start, ws.max_row + 1):
        if ws[f"{col_letter}{r}"].value == text:
            return r
    raise AssertionError(f"{text!r} not found in column {col_letter} of {ws.title}")


def num(v) -> float:
    return float(v) if v is not None else 0.0


def dia_columns(ws, dia_row: int, first_col: int) -> Dict[int, str]:
    out = {}
    c = first_col
    while True:
        v = ws.cell(dia_row, c).value
        if not isinstance(v, (int, float)):
            break
        out[int(v)] = get_column_letter(c)
        c += 1
    return out


# --------------------------------------------------------------------------- tests
def test_sheet_names(workbook):
    assert workbook.sheetnames == [
        "Summary", "14th Beam", "14th Slab", "14th Column", "14th NSW", "14th Staircase", "Extraction Log",
    ]
    assert workbook.properties.creator == "Avighna BOQ"


def test_summary_layout_and_cross_sheet_formulas(workbook, job):
    ws = workbook["Summary"]
    assert ws["A2"].value == "14th Floor Summary Of BBS"
    assert ws["A3"].value == "Sr No"
    assert ws["B3"].value == "Description"
    assert ws["C3"].value == "Dia Wise Weight Of Bars In Tons"
    assert [ws.cell(4, c).value for c in range(3, 10)] == ["8MM", "10MM", "12MM", "16MM", "20MM", "25MM", "32MM"]
    assert ws["J3"].value == "Total Wt In Tons"
    assert ws["K3"].value == "Remark"

    labels = [ws[f"B{r}"].value for r in range(5, 10)]
    assert labels == ["14th Floor Beam", "14th Floor Slab", "Column (14th Floor)", "NSW (14th Floor)",
                      "Staircase (14th Floor)"]

    # live cross-sheet formulas pointing at each tab's kg total row / 1000
    beam = workbook["14th Beam"]
    beam_kg_row = find_row(beam, "P", " Wt In Kgs")
    assert ws["C5"].value == f"='14th Beam'!S{beam_kg_row}/1000"
    assert ws["I5"].value == f"='14th Beam'!Y{beam_kg_row}/1000"
    slab = workbook["14th Slab"]
    slab_kg_row = find_row(slab, "S", " Wt In Kgs")
    assert ws["C6"].value == f"='14th Slab'!U{slab_kg_row}/1000"
    column = workbook["14th Column"]
    col_kg_row = find_row(column, "Q", "DIAMETERWISE TOTAL WEIGHT IN KG")
    assert ws["C7"].value == f"='14th Column'!V{col_kg_row}/1000"
    nsw = workbook["14th NSW"]
    assert ws["C8"].value == f"='14th NSW'!U{find_row(nsw, 'S', ' Wt In Kgs')}/1000"
    stair = workbook["14th Staircase"]
    assert ws["C9"].value == f"='14th Staircase'!U{find_row(stair, 'S', ' Wt In Kgs')}/1000"
    for r in range(5, 10):
        assert ws[f"J{r}"].value == f"=SUM(C{r}:I{r})"
    assert ws["B10"].value == "Total"
    assert ws["C10"].value == "=SUM(C5:C9)"
    assert ws["J10"].value == "=SUM(C10:I10)"
    assert ws.sheet_properties.tabColor.rgb.endswith("ED1C24")


def test_beam_tab_header_and_formulas(workbook, job):
    ws = workbook["14th Beam"]
    assert ws["A1"].value == "Name of Work: Nishika Elysium (Avighna), Worli, Mumbai"
    assert ws["A2"].value == "Contractor Name : Mehta Jaising Builders LLP"
    assert ws["A3"].value.startswith("Reinforcement For RCC Work")
    assert ws["A4"].value == "DRAWN BY :"
    assert ws["C4"].value == "MUSHARRAF"
    assert ws["G4"].value == "LD" and ws["G5"].value == 39
    assert ws["I4"].value == "Cover" and ws["I5"].value == 30
    assert ws["K4"].value == "LOCATION" and ws["L4"].value == "Beam"
    assert ws["K5"].value == "LEVEL" and ws["L5"].value == LEVEL
    assert ws["S4"].value == "Dwg No:" and ws["U4"].value == "9680/P.P/13, 13D, 13A, 13B, 201"
    assert ws["C6"].value == "Size Of Member"
    assert ws["S6"].value == "WEIGHT CALCULATIONS"
    assert ws["B7"].value == "Description"
    assert ws["J7"].value == "Dia"
    assert ws["R7"].value == "Length of one Bar"
    assert ws["S7"].value == "DIAMETERS"
    dias = dia_columns(ws, 8, 19)
    assert list(dias) == [8, 10, 12, 16, 20, 25, 32]
    assert ws["S9"].value == 0.395 and ws["W9"].value == 2.47
    assert ws["Z7"].value == "Remarks / Formula"
    assert ws.freeze_panes == "A10"
    assert ws.page_setup.orientation == "landscape"
    assert int(ws.page_setup.paperSize) == int(ws.PAPERSIZE_A3)  # A3

    # member header rows
    r_b1a = find_row(ws, "B", "B1A")
    assert ws[f"A{r_b1a}"].value == 1 and ws[f"C{r_b1a}"].value == 2
    assert (ws[f"D{r_b1a}"].value, ws[f"E{r_b1a}"].value, ws[f"F{r_b1a}"].value) == (2.45, 0.23, 0.75)
    r_b1b = find_row(ws, "B", "B1B")
    assert ws[f"A{r_b1b}"].value == 2

    # every bar row has the cut-length and weight formulas
    result_b1a = next(m for m in job.result.members if m.tag == "B1A")
    r = r_b1a + 1
    for row in result_b1a.rows:
        assert ws[f"G{r}"].value == row.description
        assert ws[f"J{r}"].value == row.dia_mm
        assert ws[f"R{r}"].value == f"=SUM(M{r}:P{r})-(J{r}*Q{r}*0.6/1000)"
        for d, c in dias.items():
            assert ws[f"{c}{r}"].value == f'=IF(J{r}=${c}$8,C{r}*L{r}*R{r}*${c}$9," - ")'
        assert ws[f"Z{r}"].value == row.formula
        r += 1

    # totals
    kg = find_row(ws, "P", " Wt In Kgs")
    assert ws[f"S{kg}"].value == f"=SUM(S10:S{kg - 1})"
    assert ws[f"P{kg + 1}"].value == " Wt In Tons" and ws[f"S{kg + 1}"].value == f"=S{kg}/1000"
    assert ws[f"P{kg + 2}"].value == "Total Wt In Tons" and ws[f"S{kg + 2}"].value == f"=SUM(S{kg + 1}:Y{kg + 1})"
    assert ws[f"P{kg + 3}"].value == "Total (kg)" and ws[f"S{kg + 3}"].value == f"=SUM(S{kg}:Y{kg})"


def test_beam_weights_recomputed_from_cells(workbook, job):
    ws = workbook["14th Beam"]
    dias = dia_columns(ws, 8, 19)
    for tag in ("B1A", "B1B"):
        res = next(m for m in job.result.members if m.tag == tag)
        r = find_row(ws, "B", tag) + 1
        for row in res.rows:
            c_members = num(ws[f"C{r}"].value)
            bars = num(ws[f"L{r}"].value)
            legs = sum(num(ws[f"{x}{r}"].value) for x in "MNOP")
            cut = legs - num(ws[f"J{r}"].value) * num(ws[f"Q{r}"].value) * 0.6 / 1000
            unit_wt = num(ws[f"{dias[int(ws[f'J{r}'].value)]}9"].value)
            weight = c_members * bars * cut * unit_wt
            assert abs(cut - row.cut_length_m) < 1e-3, (tag, row.description)
            assert abs(weight - row.weight_kg) < 0.01, (tag, row.description, weight, row.weight_kg)
            r += 1


def test_slab_style_tabs(workbook, job):
    for sheet, tag in (("14th Slab", "S3 Near C6"), ("14th NSW", "NSW1 On B36"), ("14th Staircase", "ST1")):
        ws = workbook[sheet]
        assert ws["A4"].value.startswith("Name of Work: Nishika Elysium")
        assert ws["A7"].value == "DRAWN BY :" and ws["D7"].value == "MUSHARRAF"
        assert ws["G7"].value == "LD" and ws["G8"].value == 39
        assert ws["H7"].value == "Cover"
        assert ws["I7"].value == "LOCATION" and ws["I8"].value == "LEVEL" and ws["K8"].value == LEVEL
        assert ws["U7"].value == "DRAWING REF :"
        assert ws["A9"].value == "Sr. No." and ws["B9"].value == "MEMBER" and ws["C9"].value == "NO. of Members"
        assert ws["D10"].value == "L" and ws["E10"].value == "B" and ws["F10"].value == "D"
        assert ws["H9"].value == "Dia. Of Bar"
        assert [ws.cell(9, c).value for c in range(9, 16)] == list("abcdefg")
        assert ws["Q9"].value == "Cut length Of One Bar"
        assert ws["S9"].value == "Bars per Member" and ws["T9"].value == "Total No.of Bars"
        assert ws["U9"].value == "WEIGHT CALCULATIONS" and ws["U10"].value == "DIAMETERS"
        dias = dia_columns(ws, 11, 21)
        assert list(dias) == [8, 10, 12, 16, 20, 25, 32]
        assert ws["U12"].value == 0.395
        assert ws["AB9"].value == "Remarks / Formula"
        assert ws.freeze_panes == "A13"

        res = next(m for m in job.result.members if m.tag == tag)
        hdr = find_row(ws, "B", tag)
        assert ws[f"C{hdr}"].value == res.nos
        r = hdr + 1
        for row in res.rows:
            assert ws[f"B{r}"].value == row.description
            assert ws[f"H{r}"].value == row.dia_mm
            assert ws[f"Q{r}"].value == f"=SUM(I{r}:O{r})-(H{r}*P{r}*0.6/1000)"
            assert ws[f"T{r}"].value == f"=S{r}*C{hdr}"
            for d, c in dias.items():
                assert ws[f"{c}{r}"].value == f'=IF(H{r}=${c}$11,Q{r}*T{r}*${c}$12," - ")'
            # recompute
            legs = sum(num(ws[f"{get_column_letter(x)}{r}"].value) for x in range(9, 16))
            cut = legs - num(ws[f"H{r}"].value) * num(ws[f"P{r}"].value) * 0.6 / 1000
            total_bars = num(ws[f"S{r}"].value) * num(ws[f"C{hdr}"].value)
            unit_wt = num(ws[f"{dias[row.dia_mm]}12"].value)
            assert abs(cut - row.cut_length_m) < 1e-3
            assert abs(total_bars * cut * unit_wt - row.weight_kg) < 0.01, (sheet, row.description)
            r += 1
        kg = find_row(ws, "S", " Wt In Kgs")
        assert ws[f"U{kg}"].value == f"=SUM(U13:U{kg - 1})"
        assert ws[f"S{kg + 1}"].value == " Wt In Tons" and ws[f"U{kg + 1}"].value == f"=U{kg}/1000"
        assert ws[f"S{kg + 2}"].value == "Total Wt In Tons" and ws[f"U{kg + 2}"].value == f"=SUM(U{kg + 1}:AA{kg + 1})"
        assert ws[f"S{kg + 3}"].value == "Total (kg)"


def test_column_tab(workbook, job):
    ws = workbook["14th Column"]
    assert ws["A1"].value.startswith("Name of Work")
    assert ws["A4"].value == "DRAWN BY :" and ws["G4"].value == "LD" and ws["H4"].value == "Cover"
    assert ws["G5"].value == 39 and ws["H5"].value == 40
    assert ws["I4"].value == "LOCATION" and ws["K4"].value == "Column"
    assert ws["I5"].value == "LEVEL" and ws["K5"].value == LEVEL
    assert ws["V4"].value == "DRAWING REF :"
    assert ws["A6"].value == "Sr. No." and ws["B6"].value == "MEMBER" and ws["C6"].value == "Size Of Member"
    assert (ws["C7"].value, ws["D7"].value, ws["E7"].value) == ("L", "B", "D")
    assert ws["G6"].value == "NO. of Members" and ws["H6"].value == "Dia. Of Bar"
    assert [ws.cell(6, c).value for c in range(9, 16)] == list("abcdefg")
    assert ws["Q6"].value == "Cut length In Rmt" and ws["S6"].value == "Rings Per Layer"
    assert ws["T6"].value == "Bars per Member" and ws["U6"].value == "Total No.of Bars"
    assert ws["V6"].value == "Dia Wise Wt Of Bars In Kgs As Per Unit Wt"
    dias = dia_columns(ws, 7, 22)
    assert list(dias) == [8, 10, 12, 16, 20, 25, 32]
    assert ws["V8"].value == 0.395 and ws["Y8"].value == 1.58
    assert ws["AC6"].value == "Remarks / Formula"

    res = next(m for m in job.result.members if m.tag == "C6")
    hdr = find_row(ws, "B", "C6")
    assert (ws[f"C{hdr}"].value, ws[f"D{hdr}"].value, ws[f"E{hdr}"].value) == (0.3, 0.6, 3.8)
    r = hdr + 1
    for row in res.rows:
        assert ws[f"B{r}"].value == row.description
        assert ws[f"G{r}"].value == row.members
        assert ws[f"Q{r}"].value == f"=SUM(I{r}:O{r})-(H{r}*P{r}*0.6/1000)"
        assert ws[f"U{r}"].value == f"=G{r}*T{r}"
        for d, c in dias.items():
            assert ws[f"{c}{r}"].value == f"=IF({c}$7=$H{r},$U{r}*$Q{r}*{c}$8,0)"
        legs = sum(num(ws[f"{get_column_letter(x)}{r}"].value) for x in range(9, 16))
        cut = legs - num(ws[f"H{r}"].value) * num(ws[f"P{r}"].value) * 0.6 / 1000
        total_bars = num(ws[f"G{r}"].value) * num(ws[f"T{r}"].value)
        unit_wt = num(ws[f"{dias[row.dia_mm]}8"].value)
        assert abs(cut - row.cut_length_m) < 1e-3
        assert abs(total_bars * cut * unit_wt - row.weight_kg) < 0.01, row.description
        r += 1
    # the inner tie had per_layer=2 -> rings per layer recovered from the trace
    tie_row = find_row(ws, "B", "Inner Tie T8 @100 c/c")
    assert ws[f"S{tie_row}"].value == 2

    kg = find_row(ws, "Q", "DIAMETERWISE TOTAL WEIGHT IN KG")
    assert ws[f"Q{kg - 1}"].value == "DIAMETER OF BAR" and ws[f"V{kg - 1}"].value == "8mm"
    assert ws[f"V{kg}"].value == f"=SUM(V9:V{kg - 2})"  # sums stop before the DIAMETER OF BAR label row
    assert ws[f"Q{kg + 1}"].value == "DIAMETERWISE WEIGHT IN MT" and ws[f"V{kg + 1}"].value == f"=ROUND(V{kg}/1000,3)"
    assert ws[f"Q{kg + 2}"].value == "GROSS WEIGHT IN MT" and ws[f"V{kg + 2}"].value == f"=SUM(V{kg + 1}:AB{kg + 1})"
    assert ws[f"Q{kg + 3}"].value == "Total (kg)"


def test_extraction_log(workbook, job):
    ws = workbook["Extraction Log"]
    assert [ws.cell(2, c).value for c in range(1, 11)] == [
        "Sr", "Member tag", "Kind", "Nos", "Source sheet", "Confidence", "Evidence", "Warnings", "Remarks",
        "Weight (kg)",
    ]
    tags = [ws[f"B{r}"].value for r in range(3, 9)]
    assert tags == ["B1A", "B1B", "S3 Near C6", "C6", "NSW1 On B36", "ST1"]
    assert ws["E3"].value == "9680/P.P/13" and ws["F3"].value == 0.9
    assert "B1A 230x750" in ws["G3"].value
    b1a = next(m for m in job.result.members if m.tag == "B1A")
    assert abs(ws["J3"].value - b1a.weight_kg) < 1e-6
    # the deep beam picks up the engine's SFR assumption warning
    assert "Side face reinforcement" in (ws["H3"].value or "")
    assert ws["A9"].value is None


def test_dispatch_and_reload(job, tmp_path):
    out = write_workbook(job, tmp_path / "dispatch.xlsx")
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames[0] == "Summary"
    # also loads with cached values mode without error (no values yet - formulas uncalculated)
    openpyxl.load_workbook(out, data_only=True)


def test_extra_diameter_adds_summary_column(tmp_path):
    std = Standards()
    members = [BeamSpec(tag="B40", span_m=6.0, width_m=0.4, depth_m=0.9,
                        bottom_bars=[BarGroup(count=4, dia_mm=40)], top_bars=[BarGroup(count=2, dia_mm=25)],
                        stirrups=[StirrupZone(dia_mm=10, spacing_mm=150)])]
    job = JobRecord(id="x", project_name="P", level="5th Floor", standards=std, members=members,
                    result=compute_bbs(members, std, "5th Floor"))
    out = write_bbs_workbook(job, tmp_path / "b.xlsx")
    wb = openpyxl.load_workbook(out)
    ws = wb["Summary"]
    assert [ws.cell(4, c).value for c in range(3, 11)] == ["8MM", "10MM", "12MM", "16MM", "20MM", "25MM", "32MM", "40MM"]
    beam = wb["5th Beam"]
    assert beam["Z8"].value == 40 and beam["Z9"].value == 9.86
    assert beam["AA7"].value == "Remarks / Formula"
